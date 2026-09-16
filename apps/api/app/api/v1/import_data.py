from __future__ import annotations

import csv
import io
from zipfile import BadZipFile

from fastapi import APIRouter, Depends, File, UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import get_db
from app.core.exceptions import BadRequestError
from app.core.permissions import require
from app.models.user import User
from app.schemas.import_data import ImportResult
from app.services import image_warm_service
from app.services.inventory import import_service

router = APIRouter()

#: The most an import upload may weigh and the most rows it may carry. Both are
#: read into memory whole — `upload.read()` then a full `list(reader)` — on a
#: single-worker API, so an unbounded file is a way to stall the storefront with
#: one request. 5 MB mirrors the image uploader's cap (`uploads.MAX_FILE_SIZE`),
#: and 5,000 rows is comfortably above the whole catalogue while refusing a file
#: that is a catalogue-import route by URL and a memory bomb by content.
_MAX_UPLOAD_BYTES = 5 * 1024 * 1024
_MAX_UPLOAD_ROWS = 5000


def _guard_upload_size(content: bytes) -> None:
    if len(content) > _MAX_UPLOAD_BYTES:
        raise BadRequestError(
            f"File too large. Maximum size is {_MAX_UPLOAD_BYTES // (1024 * 1024)} MB"
        )


def _guard_row_count(rows: list[dict]) -> list[dict]:
    if len(rows) > _MAX_UPLOAD_ROWS:
        raise BadRequestError(
            f"Too many rows. Maximum is {_MAX_UPLOAD_ROWS} per import."
        )
    return rows


async def _parse_csv(upload: UploadFile) -> list[dict]:
    content = await upload.read()
    _guard_upload_size(content)
    return _parse_csv_content(content)


def _parse_csv_content(content: bytes) -> list[dict]:
    text = content.decode("utf-8-sig")  # handle BOM
    reader = csv.DictReader(io.StringIO(text))
    return _guard_row_count(list(reader))


async def _parse_recipe_upload(upload: UploadFile) -> list[dict]:
    """Read the editable recipe sheet and deliberately ignore workbook references."""
    content = await upload.read()
    _guard_upload_size(content)
    filename = (upload.filename or "").lower()
    if not (filename.endswith(".xlsx") or content.startswith(b"PK")):
        return _parse_csv_content(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise BadRequestError("Recipes workbook must be a valid .xlsx file") from exc
    if "Recipes" not in workbook.sheetnames:
        workbook.close()
        raise BadRequestError(
            "Recipes workbook must include an editable 'Recipes' sheet"
        )
    # Do not enumerate the workbook: `Owner reference` is deliberately a
    # read-only lookup tab and must never be treated as import data.
    try:
        rows = workbook["Recipes"].iter_rows(values_only=True)
        try:
            headers = [str(value or "").strip() for value in next(rows)]
        except StopIteration as exc:
            raise BadRequestError("Recipes worksheet is empty") from exc
        if not any(headers):
            raise BadRequestError("Recipes worksheet has no headers")
        if len(headers) != len(set(headers)):
            raise BadRequestError("Recipes worksheet has duplicate headers")
        return _guard_row_count(
            [
                dict(zip(headers, values, strict=False))
                for values in rows
                if any(value is not None and str(value).strip() for value in values)
            ]
        )
    finally:
        workbook.close()


@router.post("/categories", response_model=ImportResult)
async def import_categories(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _admin: User = Depends(require("admin.data.manage")),
):
    """Import categories from Foodics CSV export."""
    rows = await _parse_csv(file)
    result = await import_service.import_categories(db, rows)
    image_warm_service.warm_in_background(result.image_urls)
    return result


@router.post("/products", response_model=ImportResult)
async def import_products(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _admin: User = Depends(require("admin.data.manage")),
):
    """Import products from Foodics CSV export."""
    rows = await _parse_csv(file)
    result = await import_service.import_products(db, rows)
    # A Foodics export brings in the whole catalogue at once, which is exactly
    # the case that used to leave every product cold until a customer opened it.
    image_warm_service.warm_in_background(result.image_urls)
    return result


@router.post("/modifiers", response_model=ImportResult)
async def import_modifiers(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _admin: User = Depends(require("admin.data.manage")),
):
    """Import modifiers from Foodics CSV export."""
    rows = await _parse_csv(file)
    return await import_service.import_modifiers(db, rows)


@router.post("/modifier-options", response_model=ImportResult)
async def import_modifier_options(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _admin: User = Depends(require("admin.data.manage")),
):
    """Import modifier options from Foodics CSV export."""
    rows = await _parse_csv(file)
    return await import_service.import_modifier_options(db, rows)


@router.post("/product-modifiers", response_model=ImportResult)
async def import_product_modifiers(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _admin: User = Depends(require("admin.data.manage")),
):
    """Import product-modifier assignments from Foodics CSV export."""
    rows = await _parse_csv(file)
    return await import_service.import_product_modifiers(db, rows)


@router.post("/inventory-items", response_model=ImportResult)
async def import_inventory_items(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _admin: User = Depends(require("inventory.manage")),
):
    """Import the editable inventory catalogue exported by MM."""
    return await import_service.import_inventory_items(db, await _parse_csv(file))


@router.post("/recipes", response_model=ImportResult)
async def import_recipes(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _admin: User = Depends(require("catalogue.recipes.manage")),
):
    """Stage Recipes-sheet changes as drafts; reference sheets are ignored."""
    return await import_service.import_recipes(db, await _parse_recipe_upload(file))
